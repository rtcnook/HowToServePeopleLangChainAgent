"""
Base tools for the HowToServePeopleLangChainAgent system.
"""
import os
from pathlib import Path

from dotenv import load_dotenv
from langchain_community.document_loaders import WebBaseLoader
from langchain_core.tools import tool

load_dotenv()

# ── Web Search (Tavily) ─────────────────────────────────────────────────────

from tavily import TavilyClient

_tavily_client = TavilyClient(api_key=os.getenv("TAVILY_API_KEY", ""))


@tool
def web_search(query: str) -> str:
    """Search the web for up-to-date information. Returns results with titles, URLs,
    and content snippets.  Source attribution is included."""
    resp = _tavily_client.search(query, max_results=5, include_raw_content=False)
    results = resp.get("results", [])
    if not results:
        return f"Tavily 搜索无结果。查询: {query}"
    lines = [f"搜索: {query}", ""]
    for i, r in enumerate(results, 1):
        lines.append(f"{i}. {r.get('title', '无标题')}")
        lines.append(f"   URL: {r.get('url', 'N/A')}")
        lines.append(f"   {r.get('content', '')[:400]}")
        lines.append("")
    return "\n".join(lines)


# ── URL Reader ──────────────────────────────────────────────────────────────

@tool
def read_url(url: str) -> str:
    """Read the full text content of a web page given its URL.
    Returns the extracted text, or an error if inaccessible."""
    try:
        loader = WebBaseLoader(url)
        docs = loader.load()
        if not docs:
            return f"无法加载: {url}"
        content = docs[0].page_content
        if len(content) > 12000:
            content = content[:12000] + "\n\n... [已截断]"
        return f"[网页内容] {url}\n\n{content}"
    except Exception as e:
        return f"读取 URL 失败 ({url}): {e}"


# ── Document Parser (PDF / Word) ───────────────────────────────────────────

@tool
def parse_document(file_path: str) -> str:
    """Parse a recruitment document (PDF or .docx) and extract all its text content.

    Use this when the user uploads a 招聘公告 / 职位表 / 报考简章 file.
    Supports .pdf and .docx formats.

    Returns the full text, or an error if parsing fails."""
    path = Path(file_path)
    if not path.exists():
        return f"文件不存在: {file_path}"

    suffix = path.suffix.lower()

    try:
        if suffix == ".pdf":
            return _parse_pdf(path)
        elif suffix in (".docx", ".doc"):
            return _parse_docx(path)
        else:
            return f"不支持的文件格式: {suffix}。请提供 .pdf 或 .docx 文件。"
    except Exception as e:
        return f"解析文件失败 ({file_path}): {e}"


def _parse_pdf(path: Path) -> str:
    import pymupdf

    doc = pymupdf.open(str(path))
    pages = []
    for page in doc:
        text = page.get_text()
        if text.strip():
            pages.append(text)
    doc.close()

    full = "\n\n".join(pages)
    if len(full) > 15000:
        full = full[:15000] + "\n\n... [已截断]"
    return f"[PDF] {path.name}\n\n{full}"


def _parse_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    paragraphs = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text)

    # Also try to extract tables
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))

    full = "\n".join(paragraphs)
    if len(full) > 15000:
        full = full[:15000] + "\n\n... [已截断]"
    return f"[DOCX] {path.name}\n\n{full}"
